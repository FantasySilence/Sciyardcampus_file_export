# ================================================================== #
# @Author: Fantasy_Silence                                           #
# @Time: 2026-10-29                                                  #
# @IDE: Visual Studio Code & PyCharm                                 #
# @Python: 3.9.7                                                     #
# ================================================================== #
# @Description: 把题目图片和 txt 合成为按题型分开的题库 Word           #
# ================================================================== #
import os
import re
import glob
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from typing import List, Optional, Tuple
from docx.enum.text import WD_ALIGN_PARAGRAPH

from src.common import const
from src.common.filesio import FilesIO
from src.common.figuresio import FiguresIO


class WordBuilder:

    """
    把 row_pngs 和 row_txt 合成题库 Word，按题型各存一个文件：
    resources/word_output/<题型>题库.docx

    规则（txt 的文件名决定用哪种）：
        1_formula.txt / 1_table.txt / 1_image.txt（以及 1_formula_table.txt 这类组合）
                      —— 有公式、表格或图像，OCR 认不准、转文字会丢内容，直接把原图贴进 Word；
        1.txt         —— 纯文字的题，把 txt 里的段落原样转成 Word 正文。
    """

    @staticmethod
    def usable_width(document) -> int:

        """
        页面可用宽度（EMU）= 纸宽 - 左右边距，插图按这个宽度铺满正文区
        """

        section = document.sections[0]
        return section.page_width - section.left_margin - section.right_margin

    @staticmethod
    def setup(document):

        """
        页面和字体：A4、2.54 厘米页边距、宋体五号
        """

        section = document.sections[0]
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)

        normal = document.styles["Normal"]
        normal.font.name = const.WORD_FONT
        normal.font.size = Pt(const.WORD_FONT_SIZE)
        normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), const.WORD_FONT)

    @staticmethod
    def types() -> List[str]:

        """
        row_txt 下的题型文件夹（选择题、分析题 ...）
        """

        root = FilesIO.loadText()
        return sorted(name for name in os.listdir(root) if os.path.isdir(os.path.join(root, name)))

    @staticmethod
    def indexes(png_type: str) -> List[int]:

        """
        这一题型下有哪些题号（按 1.txt / 1_formula.txt / 1_image.txt 这类文件名取数字）
        """

        root = FilesIO.loadText(sub_dir=png_type)
        indexes = set()
        for name in os.listdir(root):
            match = re.match(r"^(\d+)", name)
            if match:
                indexes.add(int(match.group(1)))
        return sorted(indexes)

    @staticmethod
    def paragraphs(path: str) -> List[str]:

        """
        读 txt：段与段之间空一行
        """

        with open(path, encoding="utf-8") as file:
            return [paragraph.strip() for paragraph in file.read().split("\n\n") if paragraph.strip()]

    @staticmethod
    def add_title(document, index: int):

        """
        每道题前面的小标题
        """

        run = document.add_paragraph().add_run(const.WORD_QUESTION_TITLE % index)
        run.bold = True

    @staticmethod
    def add_text(document, path: str):

        """
        没公式的题：txt 里的每一段转成一个 Word 段落
        """

        for paragraph in WordBuilder.paragraphs(path):
            document.add_paragraph(paragraph)

    @staticmethod
    def add_picture(document, path: str):

        """
        有公式的题：把题目原图贴进来，铺满正文宽度
        """

        document.add_picture(path, width=WordBuilder.usable_width(document))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    @staticmethod
    def build(png_type: str) -> Optional[Tuple[str, int, int]]:

        """
        生成一种题型的题库，返回 (文件路径, 文字题数, 图片题数)
        """

        indexes = WordBuilder.indexes(png_type)
        if not indexes:
            print("[提示] %s 下没有 txt，先跑 test_4.py" % png_type)
            return None

        document = Document()
        WordBuilder.setup(document)

        text_count = 0
        image_count = 0
        for index in indexes:

            txt_root = FilesIO.loadText(sub_dir=png_type)
            tagged = sorted(glob.glob(os.path.join(txt_root, "%d_*.txt" % index)))
            plain_txt = os.path.join(txt_root, "%d.txt" % index)
            png_path = os.path.join(FiguresIO.loadFigure(sub_dir=png_type), "%d.png" % index)

            # ------ 标了公式/表格/图像：贴原图 ------ #
            if tagged:
                WordBuilder.add_title(document, index)
                if os.path.exists(png_path):
                    WordBuilder.add_picture(document, png_path)
                    image_count += 1
                else:
                    print("[警告] %s\\%d.png 不存在，%s 退化成文字"
                          % (png_type, index, os.path.basename(tagged[0])))
                    WordBuilder.add_text(document, tagged[0])
                    text_count += 1
                continue

            # ------ 纯文字：转文字 ------ #
            if os.path.exists(plain_txt):
                WordBuilder.add_title(document, index)
                WordBuilder.add_text(document, plain_txt)
                text_count += 1
                continue

            print("[提示] %s 第 %d 题没有 txt，跳过" % (png_type, index))

        target = FilesIO.loadWord("%s%s.docx" % (png_type, const.WORD_TITLE_SUFFIX))
        document.save(target)
        return target, text_count, image_count

    @staticmethod
    def save(png_type: Optional[str] = None):

        """
        png_type: 题型，不传就把 row_txt 下所有题型各生成一个 docx
        """

        kinds = [png_type] if png_type is not None else WordBuilder.types()
        if not kinds:
            print("[提示] row_txt 下还没有题型文件夹，先跑 test_3.py 截图、test_4.py 提取文字")
            return

        for kind in kinds:
            result = WordBuilder.build(kind)
            if result is None:
                continue
            target, text_count, image_count = result
            print("已保存 %s（文字 %d 题，图片 %d 题）"
                  % (os.path.basename(target), text_count, image_count))
